import sys
import unittest
from datetime import datetime, timezone
from pathlib import Path
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from cemaden_drought import INDEX_URL, discover_bulletins, discover_workbook, build_product, update_data

NOW = datetime(2026, 10, 13, 13, 0, tzinfo=timezone.utc)
PAGE = INDEX_URL + '/monitoramento-de-secas-e-impactos-no-brasil-2013-setembro-2026'
FILE = PAGE + '/IIS_Brasil_2026_09.xlsx'
REGISTRY = {'1721000': 'Palmas', '1702109': 'Araguaína', '1708205': 'Formoso do Araguaia', '1700400': 'Almas'}


def workbook():
    book = Workbook()
    sheet = book.active
    sheet.title = 'IIS'
    sheet.append(['CD_MUN', 'NM_MUNICIP', 'UF', 'IIS6_2026_09', 'IIS3_2026_08', 'IIS3_2026_09', 'IIS3_2026_11'])
    for code, name, previous, value in [('1721000', 'PALMAS', 5, 6), ('1702109', 'ARAGUA�NA', 4, 5), ('1708205', 'FORMOSO', 5, 4), ('1700400', 'ALMAS', 1, 1)]:
        sheet.append([int(code), name, 'TO', 3, previous, value, 1])
    sheet.append([5200050, 'Outro estado', 'GO', 1, 1, 1, 1])
    legend = book.create_sheet('Classes de seca')
    legend.append(['Classes de seca (IIS)'])
    for label in ['1 - Seca Excepcional', '2 - Seca Extrema', '3 - Seca Severa', '4 - Seca Moderada', '5 - Seca Fraca', '6 - Normal']:
        legend.append([label])
    return book


class DroughtTests(unittest.TestCase):
    def test_discovery_uses_official_links_and_published_period(self):
        html = f'<a href="{PAGE}">Setembro</a><a href="{PAGE}">Repetido</a><a href="{INDEX_URL}/monitoramento-de-secas-e-impactos-no-brasil-2013-novembro-2026">Futuro</a><a href="https://example.test/setembro-2026">Não oficial</a>'
        editions = discover_bulletins(html, NOW)
        self.assertEqual(editions, [((2026, 9), PAGE)])
        self.assertEqual(discover_workbook(f'<a href="{FILE}/view">arquivo</a>', PAGE, (2026, 9)), FILE)
        with self.assertRaises(ValueError):
            discover_workbook(f'<a href="{PAGE}/IIS_Brasil_2026_08.xlsx">errado</a>', PAGE, (2026, 9))

    def test_official_categories_and_code_join(self):
        result = build_product(workbook(), REGISTRY, (2026, 9), FILE, PAGE, NOW)
        self.assertEqual(result['situacao_geral'], 'Excepcional')
        self.assertEqual(result['resumo']['com_seca'], 3)
        self.assertEqual(result['resumo']['sem_seca'], 1)
        self.assertEqual(result['resumo']['moderada_ou_superior'], 2)
        self.assertEqual(result['resumo']['severa_ou_extrema'], 1)
        self.assertEqual(result['referencia'], '2026-09-01T00:00:00Z')
        city = next(row for row in result['municipios'] if row['codigoIbge'] == '1702109')
        self.assertEqual(city['nome'], 'Araguaína')
        self.assertEqual(city['classe'], 'Fraca')
        self.assertEqual(city['classe_oficial'], 'Seca Fraca')
        self.assertEqual(city['tendencia'], 'Melhora')
        self.assertEqual(result['tendencia'], 'Melhora')

    def test_missing_invalid_duplicate_or_incomplete_data_is_not_normal(self):
        for value in [None, 0, 7, True, -1, '']:
            book = workbook()
            book['IIS']['F2'] = value
            with self.assertRaises(ValueError):
                build_product(book, REGISTRY, (2026, 9), FILE, PAGE, NOW)
        book = workbook()
        book['IIS'].delete_rows(2)
        with self.assertRaises(ValueError):
            build_product(book, REGISTRY, (2026, 9), FILE, PAGE, NOW)
        book = workbook()
        book['IIS'].append(list(book['IIS'].values)[1])
        with self.assertRaises(ValueError):
            build_product(book, REGISTRY, (2026, 9), FILE, PAGE, NOW)

    def test_legend_changes_require_review_and_trend_requires_previous_month(self):
        book = workbook()
        book['Classes de seca']['A2'] = '1 - Outra categoria'
        with self.assertRaises(ValueError):
            build_product(book, REGISTRY, (2026, 9), FILE, PAGE, NOW)
        book = workbook()
        book['IIS']['E2'] = None
        result = build_product(book, REGISTRY, (2026, 9), FILE, PAGE, NOW)
        self.assertIsNone(result['tendencia'])
        self.assertEqual(result['resumo']['comparados'], 3)

    def test_normal_municipalities_are_not_critical(self):
        book = workbook()
        for row in range(2, 6):
            book['IIS'].cell(row, 6, 6)
        result = build_product(book, REGISTRY, (2026, 9), FILE, PAGE, NOW)
        self.assertEqual(result['situacao_geral'], 'Sem seca')
        self.assertEqual(result['resumo']['com_seca'], 0)
        self.assertEqual(result['resumo']['municipios_criticos'], [])

    def test_update_preserves_other_sources_and_last_success_on_failure(self):
        data = {'seca': {'referencia': 'antiga'}, 'erros_atualizacao': {'s2id': '403'}, 'qualidade_fontes': {'seca_iis3': {'ultimoSucesso': 'antes'}}}
        def failed():
            raise ValueError('Arquivo indisponível')
        self.assertFalse(update_data(data, failed, NOW))
        self.assertEqual(data['seca']['referencia'], 'antiga')
        self.assertEqual(data['qualidade_fontes']['seca_iis3']['ultimoSucesso'], 'antes')
        self.assertEqual(data['qualidade_fontes']['seca_iis3']['status'], 'erro')
        self.assertEqual(data['erros_atualizacao']['s2id'], '403')
        product = build_product(workbook(), REGISTRY, (2026, 9), FILE, PAGE, NOW)
        self.assertTrue(update_data(data, lambda: product, NOW))
        self.assertNotIn('seca_iis3', data['erros_atualizacao'])
        self.assertEqual(data['qualidade_fontes']['seca_iis3']['status'], 'ok')
        self.assertEqual(data['erros_atualizacao']['s2id'], '403')


if __name__ == '__main__':
    unittest.main()
