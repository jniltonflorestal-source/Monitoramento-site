"""Collect the public monthly CEMADEN workbook; do not use the restricted map API."""
import argparse
import io
import json
import math
import os
import re
import sys
import tempfile
import time
from datetime import datetime, timezone
from html.parser import HTMLParser
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.parse import urljoin, urlparse
from urllib.request import Request, urlopen
from zipfile import ZipFile
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
INDEX_URL = 'https://www.gov.br/cemaden/pt-br/assuntos/monitoramento/impactos-seca/monitoramento-de-seca-para-o-brasil'
MONTHS = ['janeiro', 'fevereiro', 'marco', 'abril', 'maio', 'junho', 'julho', 'agosto', 'setembro', 'outubro', 'novembro', 'dezembro']
CLASSES = {1: ('Excepcional', 'Seca Excepcional'), 2: ('Extrema', 'Seca Extrema'), 3: ('Severa', 'Seca Severa'), 4: ('Moderada', 'Seca Moderada'), 5: ('Fraca', 'Seca Fraca'), 6: ('Sem seca', 'Normal')}


def stamp(date):
    return date.astimezone(timezone.utc).isoformat(timespec='seconds').replace('+00:00', 'Z')


class Links(HTMLParser):
    def __init__(self, html):
        super().__init__(convert_charrefs=True)
        self.urls = []
        self.feed(html)

    def handle_starttag(self, tag, attrs):
        if tag == 'a':
            href = dict(attrs).get('href')
            if href:
                self.urls.append(href)


def official_url(url):
    parsed = urlparse(url)
    return parsed.scheme == 'https' and parsed.netloc == 'www.gov.br' and parsed.path.startswith(urlparse(INDEX_URL).path + '/') and not parsed.query


def discover_bulletins(html, now):
    editions = set()
    for href in Links(html).urls:
        url = urljoin(INDEX_URL + '/', href).rstrip('/')
        match = re.search(r'/monitoramento-de-secas-e-impactos-no-brasil-(?:2013-)?([a-z]+)-(\d{4})$', url)
        if not official_url(url) or not match or match[1] not in MONTHS:
            continue
        period = (int(match[2]), MONTHS.index(match[1]) + 1)
        if 2000 <= period[0] and period <= (now.year, now.month):
            editions.add((period, url))
    if not editions:
        raise ValueError('Boletim mensal CEMADEN não identificado na página oficial')
    return sorted(editions, reverse=True)


def discover_workbook(html, page, period):
    filename = f'IIS_Brasil_{period[0]}_{period[1]:02d}.xlsx'
    for href in Links(html).urls:
        url = urljoin(page + '/', href).removesuffix('/view')
        if official_url(url) and urlparse(url).path.endswith('/' + filename):
            return url
    raise ValueError('Planilha IIS do período publicado não encontrada')


def fetch_bytes(url, max_bytes=8 * 1024 * 1024):
    if url != INDEX_URL and not official_url(url):
        raise ValueError('Endereço fora da publicação oficial CEMADEN')
    for attempt in range(2):
        try:
            request = Request(url, headers={'User-Agent': 'DefesaCivil-Tocantins/1.0 (public CEMADEN bulletin collector)'})
            with urlopen(request, timeout=20) as response:
                if response.url != INDEX_URL and not official_url(response.url):
                    raise ValueError('Redirecionamento fora da fonte oficial')
                payload = response.read(max_bytes + 1)
            if len(payload) > max_bytes:
                raise ValueError('Arquivo CEMADEN excede o limite de tamanho')
            return payload
        except HTTPError as error:
            if attempt or error.code not in [408, 429, 500, 502, 503, 504]:
                raise ValueError(f'Publicação CEMADEN indisponível (HTTP {error.code})') from error
        except (URLError, TimeoutError):
            if attempt:
                raise ValueError('Falha de conexão com a publicação CEMADEN')
        time.sleep(1)


def level(value):
    if value is None or isinstance(value, bool) or str(value).strip() == '':
        return None
    try:
        number = float(value)
        return int(number) if math.isfinite(number) and number.is_integer() and number in CLASSES else None
    except (TypeError, ValueError):
        return None


def build_product(book, registry, period, file_url, page, now):
    if 'IIS' not in book.sheetnames or 'Classes de seca' not in book.sheetnames:
        raise ValueError('Planilha sem dados IIS ou legenda oficial')
    legend = {}
    for row in book['Classes de seca'].iter_rows(values_only=True):
        match = re.fullmatch(r'([1-6])\s*-\s*(.+)', str(row[0] or '').strip())
        if match:
            legend[int(match[1])] = match[2].strip()
    if legend != {code: labels[1] for code, labels in CLASSES.items()}:
        raise ValueError('Legenda oficial IIS mudou; classificação requer revisão')
    sheet = book['IIS']
    if sheet.max_row > 8000 or sheet.max_column > 256:
        raise ValueError('Dimensões da planilha IIS fora do contrato esperado')
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows)
    current = f'IIS3_{period[0]}_{period[1]:02d}'
    previous_period = (period[0], period[1] - 1) if period[1] > 1 else (period[0] - 1, 12)
    previous = f'IIS3_{previous_period[0]}_{previous_period[1]:02d}'
    required = ['CD_MUN', 'NM_MUNICIP', 'UF', current]
    if any(headers.count(key) != 1 for key in required):
        raise ValueError('Colunas municipais/IIS-3 do período não confirmadas')
    columns = {key: headers.index(key) for key in required}
    previous_col = headers.index(previous) if headers.count(previous) == 1 else None
    reference = f'{period[0]}-{period[1]:02d}-01T00:00:00Z'
    municipalities = {}
    improved = worsened = stable = compared = 0
    for row in rows:
        if row[columns['UF']] != 'TO':
            continue
        code = str(row[columns['CD_MUN']]).strip()
        value = level(row[columns[current]])
        if code not in registry or code in municipalities or value is None:
            raise ValueError('Município duplicado, não identificado ou categoria IIS ausente/inválida')
        old = level(row[previous_col]) if previous_col is not None else None
        trend = None
        if old is not None:
            compared += 1
            trend = 'Melhora' if value > old else 'Agravamento' if value < old else 'Estabilidade'
            improved += trend == 'Melhora'
            worsened += trend == 'Agravamento'
            stable += trend == 'Estabilidade'
        municipalities[code] = {'codigoIbge': code, 'nome': registry[code], 'nome_fonte': row[columns['NM_MUNICIP']], 'nivel': value, 'classe': CLASSES[value][0], 'classe_oficial': CLASSES[value][1], 'referencia': reference, 'tendencia': trend}
    if set(municipalities) != set(registry):
        raise ValueError('Cobertura municipal incompleta; não confirmado o conjunto do Tocantins')
    values = list(municipalities.values())
    worst = min(row['nivel'] for row in values)
    complete_trend = compared == len(registry)
    trend = ('Melhora' if improved > worsened else 'Agravamento' if worsened > improved else 'Estabilidade') if complete_trend else None
    return {
        'produto': 'IIS3',
        'fonte': 'CEMADEN/MCTI · IIS-3 municipal',
        'fonte_url': page,
        'arquivo_url': file_url,
        'atualizado_em': stamp(now),
        'referencia': reference,
        'referencia_anterior': f'{previous_period[0]}-{previous_period[1]:02d}-01T00:00:00Z' if previous_col is not None else None,
        'situacao_geral': CLASSES[worst][0],
        'tendencia': trend,
        'cobertura_completa': True,
        'metodo': 'Categorias oficiais do IIS-3 publicadas no boletim mensal; não é alerta ou declaração de emergência.',
        'resumo': {
            'total_municipios': len(values),
            'com_seca': sum(row['nivel'] <= 5 for row in values),
            'sem_seca': sum(row['nivel'] == 6 for row in values),
            'moderada_ou_superior': sum(row['nivel'] <= 4 for row in values),
            'severa_ou_extrema': sum(row['nivel'] <= 3 for row in values),
            'municipios_criticos': [row['nome'] for row in values if worst < 6 and row['nivel'] == worst],
            'melhoraram': improved if complete_trend else None,
            'agravaram': worsened if complete_trend else None,
            'estaveis': stable if complete_trend else None,
            'comparados': compared
        },
        'municipios': values
    }


def collect_drought():
    now = datetime.now(timezone.utc)
    period, page = discover_bulletins(fetch_bytes(INDEX_URL).decode('utf-8'), now)[0]
    file_url = discover_workbook(fetch_bytes(page).decode('utf-8'), page, period)
    payload = fetch_bytes(file_url)
    with ZipFile(io.BytesIO(payload)) as archive:
        if sum(item.file_size for item in archive.infolist()) > 128 * 1024 * 1024:
            raise ValueError('Planilha descompactada excede limite')
    geometry = json.loads((ROOT / 'data/tocantins_municipios.geojson').read_text(encoding='utf-8'))
    registry = {str(feature['properties']['codarea']): feature['properties']['nome'] for feature in geometry['features']}
    if len(registry) != 139:
        raise ValueError('Malha municipal do Tocantins não confirmada')
    book = load_workbook(io.BytesIO(payload), read_only=True, data_only=True, keep_links=False)
    try:
        return build_product(book, registry, period, file_url, page, datetime.now(timezone.utc))
    finally:
        book.close()


def update_data(data, collector=collect_drought, now=None):
    attempted = stamp(now or datetime.now(timezone.utc))
    quality = data.setdefault('qualidade_fontes', {})
    errors = data.setdefault('erros_atualizacao', {})
    try:
        product = collector()
        data['seca'] = product
        quality['seca_iis3'] = {'status': 'ok', 'ultimaTentativa': attempted, 'ultimoSucesso': product['atualizado_em']}
        errors.pop('seca_iis3', None)
        data.setdefault('fontes', {})['seca'] = product['fonte']
        data.setdefault('automacao', {})['seca_iis3'] = 'automatico_boletim_mensal_cemaden'
        data['atualizado_em'] = product['atualizado_em']
        return True
    except Exception as error:
        message = str(error)
        quality['seca_iis3'] = {**quality.get('seca_iis3', {}), 'status': 'erro', 'ultimaTentativa': attempted, 'erro': message}
        errors['seca_iis3'] = message
        return False


def main():
    sys.stdout.reconfigure(encoding='utf-8')
    parser = argparse.ArgumentParser()
    parser.add_argument('--preview', action='store_true', help='Print the official product without changing operational JSON')
    args = parser.parse_args()
    if args.preview:
        print(json.dumps(collect_drought(), ensure_ascii=False, indent=2))
        return
    target = ROOT / 'dados-monitoramento.json'
    data = json.loads(target.read_text(encoding='utf-8'))
    success = update_data(data)
    with tempfile.NamedTemporaryFile(mode='w', encoding='utf-8', dir=target.parent, delete=False, suffix='.json') as temporary:
        json.dump(data, temporary, ensure_ascii=False, indent=2)
        temporary.write('\n')
    os.replace(temporary.name, target)
    print('Seca CEMADEN atualizada' if success else 'Seca CEMADEN indisponível; último produto preservado')


if __name__ == '__main__':
    main()
